#!/usr/bin/env python3
"""Build Word document from exam_guide_ru_blocks.txt (run from repo root or this dir)."""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_LINE_SPACING
except ImportError:
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent / ".docx-deps"))
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_LINE_SPACING


def main() -> None:
    base = Path(__file__).resolve().parent
    src = base / "exam_guide_ru_blocks.txt"
    out = base / "Claude-Certified-Architect-Foundations-Exam-Guide-RU.docx"

    raw = src.read_text(encoding="utf-8")
    lines = raw.splitlines()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    paragraph_format = style.paragraph_format
    paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("===") and line.endswith("==="):
            kind = line.strip("=").strip()
            i += 1
            buf: list[str] = []
            while i < len(lines):
                nxt = lines[i]
                if nxt.strip().startswith("===") and nxt.strip().endswith("==="):
                    break
                buf.append(nxt)
                i += 1
            text = "\n".join(buf).strip()
            if not text:
                continue
            if kind == "TITLE":
                doc.add_heading(text, level=0)
            elif kind == "H1":
                doc.add_heading(text, level=1)
            elif kind == "H2":
                doc.add_heading(text, level=2)
            elif kind == "H3":
                doc.add_heading(text, level=3)
            elif kind == "P":
                doc.add_paragraph(text)
            elif kind == "BULLET":
                doc.add_paragraph(text, style="List Bullet")
            else:
                doc.add_paragraph(text)
            continue
        i += 1

    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
